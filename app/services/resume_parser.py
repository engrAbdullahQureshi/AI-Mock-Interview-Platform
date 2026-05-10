import io
import logging
from typing import Optional

logger = logging.getLogger(__name__)


async def extract_text_from_pdf(file_bytes: bytes) -> Optional[str]:
    """
    Extracts all text from a PDF file.

    USES: PyPDF2 (already in requirements.txt)

    FLOW:
    1. Load PDF from bytes
    2. Loop through every page
    3. Extract and concatenate text
    4. Return combined text string
    """
    try:
        import PyPDF2

        pdf_reader = PyPDF2.PdfReader(io.BytesIO(file_bytes))
        text_parts = []

        for page_num, page in enumerate(pdf_reader.pages):
            try:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
            except Exception as e:
                logger.warning(f"Could not extract text from page {page_num}: {e}")
                continue

        full_text = "\n".join(text_parts).strip()

        if not full_text:
            logger.warning("PDF parsed but no text extracted (may be image-based)")
            return None

        logger.info(f"PDF extracted: {len(full_text)} characters from {len(pdf_reader.pages)} pages")
        return full_text

    except Exception as e:
        logger.error(f"PDF extraction failed: {e}")
        return None


async def extract_text_from_docx(file_bytes: bytes) -> Optional[str]:
    """
    Extracts all text from a DOCX (Word) file.

    USES: python-docx
    NOTE: python-docx is not in requirements.txt yet — we handle ImportError gracefully.

    FLOW:
    1. Load DOCX from bytes
    2. Loop through all paragraphs
    3. Return combined text
    """
    try:
        from docx import Document

        doc = Document(io.BytesIO(file_bytes))
        text_parts = []

        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text_parts.append(paragraph.text)

        # Also extract text from tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text_parts.append(cell.text)

        full_text = "\n".join(text_parts).strip()

        if not full_text:
            logger.warning("DOCX parsed but no text extracted")
            return None

        logger.info(f"DOCX extracted: {len(full_text)} characters")
        return full_text

    except ImportError:
        logger.error("python-docx not installed. Run: pip install python-docx")
        return None
    except Exception as e:
        logger.error(f"DOCX extraction failed: {e}")
        return None


async def extract_text_from_txt(file_bytes: bytes) -> Optional[str]:
    """Extracts text from a plain .txt file."""
    try:
        text = file_bytes.decode("utf-8").strip()
        logger.info(f"TXT extracted: {len(text)} characters")
        return text
    except UnicodeDecodeError:
        try:
            text = file_bytes.decode("latin-1").strip()
            return text
        except Exception as e:
            logger.error(f"TXT extraction failed: {e}")
            return None


async def extract_text(file_bytes: bytes, content_type: str, filename: str) -> Optional[str]:
    """
    Router function — picks the right extractor based on file type.

    SUPPORTED:
    - application/pdf              → PDF extractor
    - application/vnd.openxmlformats (docx) → DOCX extractor
    - text/plain                   → TXT extractor

    Returns raw text string or None if extraction fails.
    """
    filename_lower = filename.lower()

    if content_type == "application/pdf" or filename_lower.endswith(".pdf"):
        return await extract_text_from_pdf(file_bytes)

    elif (
        content_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        or filename_lower.endswith(".docx")
    ):
        return await extract_text_from_docx(file_bytes)

    elif content_type == "text/plain" or filename_lower.endswith(".txt"):
        return await extract_text_from_txt(file_bytes)

    else:
        logger.warning(f"Unsupported file type: {content_type} / {filename}")
        return None